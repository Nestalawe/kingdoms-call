#!/usr/bin/env python3
"""build_unit_workbook_0927.py — builds the Army Units workbook from /tmp/units_raw.json
(extracted from the live GM portal by claude/claude_unitexport_0927.js).

Every number in the sheets comes from the JSON, i.e. from the engine's own tables. The prose in
ABILITY is hand-written against the engine source and cites the magnitude the code actually uses.
No formulas anywhere: the designer sorts and reorders rows, and a formula would follow the wrong row.
"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

RAW = json.load(open('/tmp/units_raw.json'))
STAMP = RAW['stamp']

FONT = 'Arial'
HDR_FILL   = PatternFill('solid', fgColor='2F3E2F')
HDR_FONT   = Font(name=FONT, bold=True, color='FFFFFF', size=10)
EDIT_FILL  = PatternFill('solid', fgColor='FFF2CC')   # the column the designer writes in
GRP_FILL   = PatternFill('solid', fgColor='EDE7DA')
CELL_FONT  = Font(name=FONT, size=10)
NAME_FONT  = Font(name=FONT, size=10, bold=True)
THIN = Side(style='thin', color='CCCCCC')
BOX  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ─────────────────────────────────────────────────────────────────────────────────────────────
# ABILITY PROSE — written against the engine source, 2026-09-27 (build 2026.09.27-1625).
# Each entry: (ability name, full description of exactly how it works, where it hooks in)
# "hook" names the engine site so a future change can be traced back to the right function.
# ─────────────────────────────────────────────────────────────────────────────────────────────
ABIL = {
 # ── HUMAN ───────────────────────────────────────────────────────────────────────────────────
 'Shieldwall Levies': ('Shieldwall',
   'Each Shieldwall Levies unit fights at ×1.25 (+25% strength) whenever the same force contains '
   'AT LEAST TWO Shieldwall Levies units. The bonus does not grow past two — three or ten units all '
   'get the same ×1.25 — and it is counted per unit, so every levy unit in the force gets it. It '
   'applies on attack and defence alike, in every terrain, and stacks multiplicatively with terrain, '
   'quality, kinship and the commander\'s Tactical bonus.',
   'unitAbilityMult'),
 'Free Lances': ('Ride Them Down',
   'When your side WINS a battle, each Free Lances unit in the winning force adds +6% to the chance '
   'that each of the LOSER\'s units is destroyed, to a maximum of +18% (i.e. three units are the cap). '
   'It works ONLY in open country — plain and desert; in forest, jungle or mountain the bonus is 0 and '
   'the report says the Lances found no room to ride. The War-Horn of the Plains (magic item) removes '
   'the terrain condition entirely and grants the full +18% anywhere. Fully symmetric: neutral province '
   'defenders can field Free Lances and theirs bites on your force the same way. This is a CASUALTY '
   'modifier, not a strength modifier — it does nothing to who wins, only to how badly the loser is cut up.',
   'rideThemDownBonus'),
 'Yeoman Longbows': ('Massed Volley',
   'In the pre-battle archery phase, a Yeoman Longbows unit deals +10% missile damage for every OTHER '
   'missile unit in the same force, capped at +30% (three others). "Missile unit" means anything in the '
   'master missile set — Yeoman Longbows, Elven Bowmen, Shadow Archers, Ironwood Crossbows, Mountain Guns, '
   'Goblin Horde — so mixed archer lines feed the bonus. It multiplies the shot\'s damage only; it does '
   'not change how many shots are fired or the per-target softening cap.',
   'missileDamageMult'),
 'Gilded Lances': ('Shock of the Charge',
   'Each Gilded Lances unit fights at ×1.40 (+40% strength) when its side is the ATTACKER. Defending, it '
   'gets nothing at all. Terrain is unchanged by this ability, so heavy horse still suffer the ordinary '
   'forest (×0.75) and jungle (×0.70) penalties — a charge into woods is still a bad charge. Unconditional '
   'on attack: no terrain, no minimum count, no enemy condition.',
   'unitAbilityMult'),
 'Human Garrison': ('— (no combat ability)',
   'A garrison unit has no special ability. What makes it a garrison is where it goes: recruiting one always '
   'musters it straight into the province garrison rather than into a commander\'s column, it moves at speed 1, '
   'and it costs 1g to raise and 1g a turn to keep — the cheapest unit in the game. Base strength 2, the lowest '
   'of any humanoid unit. Every garrisoned unit is drilled at the end of each turn as if by a commander of '
   'Tactical 1 (Naval 1 for hulls), so garrisons do slowly improve in quality on their own.',
   'DISC_CLS_SETS.garrison / isGarrisonUnitType'),
 'Noble Griffons': ('Screech Dive (pre-battle strike)',
   'A creature unit, not a humanoid one: it needs Druid 3 to lead without penalty, and a Druidic realm or a '
   'Druidic commander waives that requirement entirely. It carries no always-on strength multiplier. Instead, '
   'in the pre-battle MAGIC phase it fires Screech Dive at ONE enemy unit for 2–4 strength damage. The chance '
   'it fires at all is min(90%, 30% + 10% per Druid level of the commanding character + any primalStrike item '
   'bonus) — so a commander with no Druid skill still gets 30%, and Druid 6 reaches the 90% ceiling. Flying: '
   'terrain never affects it, and it can lift one foot unit across water.',
   'BATTLE_ARMY_SPELLS / primalStrikeChance'),
 'Galleon Fleet': ('Broadside',
   'In a SEA battle only, each Galleon Fleet unit is added to the list of units that loose in the pre-battle '
   'archery phase — a galleon fires as a floating battery alongside whatever archers the force is carrying. It is '
   'ONE shot per galleon per volley, not a doubled shot: the galleon is not in the master missile set, so the '
   'ordinary filter skips it and this push is its only entry. On land it does not shoot at all, and it never counts '
   'toward another unit\'s Massed Volley. No damage multiplier of its own. Carrier: 1 hull of lift (2 light-foot '
   'units, or 1 heavy/cavalry/siege/beast). The Human home fleet.',
   'abilityMissileUnits'),
 'Corsair Fleet': ('Cut the Line',
   'If AT LEAST HALF the units in your winning force are Corsair Fleet, the beaten enemy may not retreat by water: '
   'every sea province is struck off their retreat list, and if that leaves them nothing they are treated as trapped. '
   'Against NEUTRAL province guards — who never withdraw, so there is no line to sever — it inverts into Keep the Sea '
   'Lanes: if your own beaten force stands on ground adjacent to any sea, the trapped-in-place casualty roll is SUPPRESSED. Be clear '
   'about what that does and does not do — it cancels the losses, but nothing actually moves: the force stays where it is and the '
   'province is not left contested. It fires only when there is no friendly retreat at all, and Deathless March and Shield Fort are '
   'both checked ahead of it. Carrier: 1 hull of lift. NOTE: the Corsair Fleet is retired '
   'from world generation — coastal Human realms now build Galleon Fleet only — so it appears in play only where an '
   'existing game already holds one.',
   'isHalfOrMore / PvP retreat block / combatVsNPC'),
 # ── ELVEN ───────────────────────────────────────────────────────────────────────────────────
 'Elven Bowmen': ('Woodland Volley',
   'DOUBLE missile damage (×2) in forest and jungle. Nothing in plain, desert, mountain or at sea. This is the largest '
   'single missile multiplier in the game and it stacks with Massed Volley, so elven bows massed in their own woods are '
   'brutal. An Elven unit of an Elven realm also carries Elven Woodland Mastery (×1.5 strength in forest and jungle) on '
   'top — a separate racial rule, not part of this ability.',
   'missileDamageMult'),
 'Forest Wardens': ('Ambush',
   'In forest or jungle, a force facing Forest Wardens loses its archery phase entirely — its missile units do not fire at all, '
   'not even at reduced effect. One Warden unit is enough; there is no half-the-force condition, and it shields the whole force, '
   'not just the Wardens. Nothing outside forest and jungle. It works the same on ATTACK as on defence, and against a neutral '
   'garrison exactly as against another player: the test is made against the other side\'s force whichever way round it is, so '
   'Wardens marching on a wood silence the volley of whoever holds it. (Until 2026-09-27 only the attacker\'s volley was gated in '
   'battles against neutral defenders, so attacking Wardens silenced a player but not a neutral garrison. Both paths now agree.)',
   'archeryBlockedBy'),
 'Shadow Archers': ('Loosed From Cover',
   'Shadow Archers are struck out of the enemy\'s missile target pool BY TYPE, before the arrow-cap test is applied. They '
   'become targetable only when they are the ONLY units left in the force — other units reaching their arrow cap does NOT '
   'expose them; it just means the enemy has nothing left to shoot at and the volley is wasted. They shoot normally '
   'themselves (they are in the master missile set), so they are pure ranged value with no return fire. Base strength 6, '
   'the highest of any elven land unit.',
   'missileTargetPool'),
 'Elven Garrison': ('— (no combat ability)',
   'No special ability; a garrison unit (see Human Garrison). The Elven garrison is the best of the four in woods — ×1.10 in '
   'forest and jungle, where the others sit at ×1.00 — and the worst in desert (×0.85).',
   'DISC_CLS_SETS.garrison'),
 'Eagle Riders': ('Skyward Retreat',
   'When a beaten force has NO line of retreat, every other unit rolls a 15% chance of being cut down in the rout. Eagle '
   'Riders skip that roll entirely — they climb out over the encircling lines, always, with no chance of loss. It applies in '
   'both player-vs-player battles and battles against neutral defenders. Flying: terrain never affects them (always ×1.00) '
   'and each Eagle Riders unit can lift ONE foot unit across water, so enough of them carry a whole column over the sea with '
   'no fleet at all. Needs Druid 3 to lead without penalty (waived for a Druidic realm/commander, or with the Dragon Taming '
   'discovery).',
   'PvP _noEscapeRoll / combatVsNPC rout block'),
 'Elven Galleys': ('Sea Archers',
   'While an Elven Galleys unit is anywhere in the force, MISSILE units in that force ignore the land-unit sea penalty: their '
   '×0.75 for fighting at sea is lifted back to ×1.00. It applies to any missile unit, not only elven ones, and to the missile '
   'units only — ordinary foot still fights at ×0.75. One galley is enough. Carrier: 1 hull of lift. The Elven home fleet.',
   'armStr sea-penalty block'),
 'Whale Cohort': ('Breach and Swamp',
   'Before either side\'s strength is counted, EACH Whale Cohort unit rolls once: on a 25% chance it picks a random enemy sea '
   'unit and sinks it outright — the unit is destroyed, with no combat, no saving roll and no arrows spent. Three cohorts roll '
   'three times. It only ever targets sea units (hulls and sea beasts), so against a purely land force it does nothing. A '
   'creature unit needing Druid 3 (waived for Druidic realms/commanders). It is a SWIMMER, not a carrier in the ordinary sense: '
   'it counts as 1 hull of lift.',
   'applyBreachAndSwamp'),
 # ── DWARVEN ─────────────────────────────────────────────────────────────────────────────────
 'Dwarven Guard': ('Shield Fort',
   'If AT LEAST HALF the units in a beaten force are Dwarven Guard, that force never routs. Trapped with no retreat, it locks '
   'shields where it stands: NO flight casualties at all, and the province is left contested (_battleUnresolved), so the battle '
   'carries on next phase instead of the ground changing hands. With a retreat open it conducts a fighting withdrawal in good order: '
   'immune to the enemy\'s Ride Them Down casualty bonus, and a unit that fails its demotion roll falls back behind the shield wall '
   'instead of losing a quality tier — or, for a Green unit, instead of disbanding, which is what dropping below Green means. '
   'Both halves work in EVERY battle, against another player and against neutral province defenders alike. It also runs for a '
   'Guard-heavy NEUTRAL garrison: such a garrison gives your Free Lances no broken backs to ride down and takes no '
   'pinned-against-their-own-ground losses (it never retreats, so the good-order half cannot apply to it). (Until 2026-09-27 the '
   'good-order half was wired in player-vs-player battles only, so against a neutral garrison a beaten Guard-heavy force still took '
   'Ride Them Down losses and still demoted, under a line saying it had withdrawn in good order.) This ability replaced the older '
   '"Hold the Line", which pinned the force to a lost field — it no longer does that. A Dwarven unit of a Dwarven realm also defends '
   'at ×1.5 (Dwarven Bulwark), a separate racial rule.',
   'isHalfOrMore / retreat branches / combatVsNPC loser-casualty block'),
 'Ironwood Crossbows': ('Punch Through',
   'Archery can only soften a target so far. The ordinary per-unit cap is 40% of the target\'s base strength (80% for a winged '
   'rider). A shot fired by Ironwood Crossbows raises that ceiling to 55% on the ground and 90% against a winged rider — so they '
   'can keep grinding a target that ordinary bows have already finished with. The CEILING is set by whoever is firing, but the damage '
   'ALREADY DONE is one shared per-target total — two Ironwood units share a single 55% pool, they do not get 55% each. So mixing '
   'crossbows with longbows lets the crossbows keep shooting at a target the longbows have maxed out, but stacking crossbows does not '
   'raise the ceiling any further. No damage multiplier of its own.',
   'missileCap'),
 'Mountain Guns': ('Bombardment',
   'Two effects. (1) Mountain Guns are sorted to the FRONT of the archery order, so they fire before every other missile unit and '
   'get the clear shot before any per-target cap fills up. (2) ×1.5 missile damage against any GARRISON unit (any unit whose type '
   'name contains "Garrison"). Base strength 7 — the strongest humanoid land unit in the game — but speed 1, the slowest, and '
   'penalised in EVERY terrain except mountain: plain ×0.85, desert ×0.85, forest ×0.80, jungle ×0.75, mountain ×1.20. For sea '
   'lift they count as light foot (half a hull), so an Ironclad can ferry a Gun and a Guard together.',
   'abilityMissileUnits / missileDamageMult'),
 'Dwarf Garrison': ('— (no combat ability)',
   'No special ability; a garrison unit (see Human Garrison). Slightly weaker than the Human garrison in open ground (×0.90 plain, '
   '×0.85 desert) and the best of the four in mountain.',
   'DISC_CLS_SETS.garrison'),
 'Gnomish Gyrocopters': ('Flak Screen',
   'Before the battle proper, each Gyrocopter unit picks ONE enemy flying unit at random and harries it for −30% of that unit\'s base '
   'strength (minimum −1), applied as temporary strength reduction for this battle. A flyer already engaged by one gyrocopter cannot '
   'be picked again in the same battle, so N gyrocopters tie up N DISTINCT flyers — doubling up on one target is impossible. (Until '
   '2026-09-27 that "already engaged" mark was never cleared, so a flyer harried once was immune to Flak Screen for the rest of the '
   'GAME; fixed in this build — the mark now dies with the battle.) It only '
   'ever targets the four winged-rider types (Eagle Riders, Noble Griffons, Gnomish Gyrocopters, Wyvern Riders), so it does nothing '
   'against primal winged beasts like Elder Dragons. A machine, not a beast: Druid 2 to lead, and the Dragon Taming discovery '
   'deliberately does NOT waive it.',
   'applyFlakScreen'),
 'Dwarf Ironclads': ('Armoured Hulls',
   'Completely immune to pre-battle missile damage. The engine treats an Ironclad as having already reached its arrow cap, so the '
   'missile target pool skips it and no shot is ever aimed at it — it is not "reduced" damage, it is zero, whoever is firing and '
   'however many. Base strength 7, the strongest hull in the game, and it fights at full strength on land as well as at sea. Carrier: '
   '1 hull of lift, doubled to 2 by a commander with Naval 8. The Dwarven home fleet.',
   'missileCapReached'),
 # ── ORCISH ──────────────────────────────────────────────────────────────────────────────────
 'Warg Riders': ('Run Them Down',
   'If AT LEAST HALF the units in your winning force are Warg Riders, your side takes up the pursuit of the beaten enemy commander '
   'AUTOMATICALLY, without spending an order on it — provided the beaten commander is still ALIVE, is not one of your own, and you do '
   'not already have a pursuit running. Against NEUTRAL province guards — who do not flee the province and so cannot be pursued — it '
   'converts into speed instead: your own casualty chance for the battle is HALVED (×0.5 on each of your units\' destruction rolls). '
   'Nothing happens when you lose.',
   'isHalfOrMore / PvP pursuit / combatVsNPC winDeathChance'),
 'Orc Berserkers': ('Outnumbered and Uncaring',
   'Each Berserker unit gains +10% strength for every FULL 25% by which the enemy force outnumbers yours, capped at +30% (×1.30). The '
   'count is of UNITS, not strength: the excess is (enemy count − your count) / your count, floored to whole 25% steps. So 5 units '
   'against 4 gives 25% excess → +10%; 7 against 4 gives 75% → +30%; 8 against 4 or worse stays at +30%. Nothing at all when you are '
   'even or ahead in numbers.',
   'unitAbilityMult'),
 'Goblin Horde': ('Sheer Numbers',
   'Each Goblin Horde unit gains +10% strength for every OTHER Goblin Horde in the same force, capped at +30% (four hordes together). '
   'Base strength 3 — the weakest unit in the game bar garrisons — so even at the cap a horde is worth less than most single units; the '
   'point is that they cost 2g. Also a MISSILE unit: hordes fire in the archery phase and count toward other units\' Massed Volley.',
   'unitAbilityMult'),
 'Orc Garrison': ('— (no combat ability)',
   'No special ability; a garrison unit (see Human Garrison). Matches the Human garrison in most terrain (×1.00 plain, forest, jungle) '
   'and takes ×0.90 in desert.',
   'DISC_CLS_SETS.garrison'),
 'Wyvern Riders': ('Terror From Above + Acid Spit Volley',
   'Two effects. (1) Terror From Above: if Wyvern Riders are present on your side, the ENEMY\'s chance of their force turning fanatical '
   'is cut by a flat 15 percentage points (floored at 0). Against neutral province guards, who never roll for fanaticism, it instead cuts '
   'the defenders\' fortification bonus by 10% (×0.90, floored at ×1.00). (2) Acid Spit Volley: a pre-battle strike on ONE enemy unit for '
   '2–5 strength damage, firing at min(90%, 30% + 10% per Druid level of the commander + primalStrike item bonus). Flying: terrain never '
   'affects them, and each lifts one foot unit over water. Druid 3 to lead, and a Druidic realm or Druidic commander does NOT waive it — '
   'that waiver covers CREATURE units only, and the Wyvern is a Monster. Only the Dragon Taming discovery removes the penalty.',
   'PvP fanaticism block / combatVsNPC defBonus / BATTLE_ARMY_SPELLS'),
 'Orcish Longships': ('Reaver Crews',
   'In the phase a force LANDS from Orcish Longships, the WHOLE force fights at ×1.25 — not just the Orcish units, and not just the ships. '
   'The bonus survives a change of command: if another hero takes charge of the combined host on the beach, the amphibious flag stays with '
   'the force and the ×1.25 still applies. It is the only unit ability in the game that lifts every unit on your side at once. Two limits worth '
   'knowing: the amphibious flag is read for the ATTACKING side only, so a force that lands and is then attacked in that same phase gets nothing; '
   'and the flag sits on the CHARACTER for the phase, so it also lifts any other battle that character starts in the same phase, not just the '
   'landing fight. Carrier: 1 hull of lift. The Orcish home fleet.',
   'unitAbilityMult (ctx.amphibious)'),
}

# Primal units: the engine's own desc already states the ability, so the prose below adds the
# mechanical detail (exact trigger, the strike chance formula, lift and Druid rules) without
# contradicting it.
PRIMAL_EXTRA = {
 'Stone Giants':     ('Mountain Bulwark',     'mountain terrain only'),
 'Forest Ents':      ('Rooted Grove',         'forest or jungle'),
 'Giant Eagles':     ('Stoop on the Archers', 'the enemy fields at least one missile unit'),
 'Cave Trolls':      ('Regenerating Hide',    'always, and again in mountain or forest'),
 'Manticore Pride':  ('Relentless Hunt',      'your side is the attacker'),
 'Basilisk Brood':   ('Stone Deterrent',      'your side is defending'),
 'Sea Giants':       ('Tidal Might',          'a COASTAL LAND province (not at sea)'),
 'Kraken Tentacles': ('Deepwater Grasp',      'a sea battle'),
 'Elder Dragons':    ('Wyrmscale',            'always — no condition'),
 'Crimson Rocs':     ('Air Superiority',      'the enemy fields at least one missile unit'),
 'Titan Warbeasts':  ('Colossal Bulk',        'always — no condition'),
 'Leviathan Pods':   ('Fleet-Sinker',         'a sea battle'),
}

ELEMENTALS = {
 'Air Elemental':      'the Howling Gale',
 'Fire Elemental':     'the Living Conflagration',
 'Earth Elemental':    'the Walking Mountain',
 'Water Elemental':    'the Surging Tide',
 'Darkness Elemental': 'the Devouring Shade',
}

RACE_ORDER = {'Human':1,'Elven':2,'Dwarven':3,'Orcish':4,'Any':5,'Primal':5,'Elemental':6}

# ARMY_RACE_OF_TYPE is the KINSHIP race (the +20% own-race / -20% feud rules) and deliberately omits
# the creature and fleet units, whose kinship race is Primal. Which REALM can raise them is a
# separate question, answered by PROVINCE_ARMIES_P for most units and by these four cases for the
# rest. Kept explicit rather than guessed, with the engine site named.
RAISED_BY_OVERRIDE = {
  'Galleon Fleet':  ('Human',  'A coastal Human realm’s home fleet (HOME_FLEET_BY_RACE / FLEET_BY_RACE) — built in a coastal or sea province, not in the land recruitment table'),
  'Corsair Fleet':  ('Human',  'RETIRED from world generation — coastal Humans now build Galleon Fleet only. Appears only in games that already hold one'),
  'Noble Griffons': ('Human',  'plain, desert, mountain (Human realms)'),
  'Whale Cohort':   ('Any',    'A wild SEA province that offers it — like a primal beast, any realm may raise it with Druid 3'),
}

def is_light(t):
    return t in {'Shieldwall Levies','Yeoman Longbows','Ironwood Crossbows','Elven Bowmen','Shadow Archers',
                 'Human Garrison','Elven Garrison','Dwarf Garrison','Orc Garrison','Goblin Horde',
                 'Forest Wardens','Dwarven Guard','Orc Berserkers','Mountain Guns'}

CARRIERS = {'Galleon Fleet','Corsair Fleet','Elven Galleys','Dwarf Ironclads','Orcish Longships','Whale Cohort'}

def build_rows():
    cost, speed, bstr = RAW['cost'], RAW['speed'], RAW['baseStr']
    terr, druid = RAW['terrain'], RAW['druidOf']
    missile  = set(RAW['missile'])
    sea      = set(RAW['seaTypes'])
    aerial   = set(RAW['aerialTypes'])
    flying   = set(RAW['flyingTypes'])
    raceof   = RAW['raceOf']
    garr     = set(RAW['classes'].get('garrison', []))
    cav      = set(RAW['classes'].get('cavalry', []))
    inf      = set(RAW['classes'].get('infantry', []))
    tamed    = set(RAW['tamedFlyers'])
    primal   = RAW['primalDefs']
    pa       = RAW['provinceArmies']

    # where each type can be recruited: terrain(s) whose race list names it
    recruit = {}
    for terrain, byrace in pa.items():
        for race, lst in byrace.items():
            for t in lst:
                recruit.setdefault(t, set()).add(terrain)

    types = sorted(set(list(cost) + list(bstr)))
    rows = []
    for t in types:
        # kinship race — the +20% own-race / -20% feud race, straight from ARMY_RACE_OF_TYPE
        if t in ELEMENTALS:
            kin = 'Elemental'
        else:
            kin = raceof.get(t, 'Primal')
        # which realm can actually raise it
        if t in RAISED_BY_OVERRIDE:
            race = RAISED_BY_OVERRIDE[t][0]
        elif t in ELEMENTALS:
            race = 'Elemental'
        elif t in primal:
            race = 'Any'
        else:
            race = raceof.get(t, 'Any')

        # class label
        if t in garr:         cls = 'Garrison'
        elif t in sea and t in primal:  cls = 'Sea beast'
        elif t in sea:        cls = 'Fleet'
        elif t in flying:     cls = 'Winged rider'
        elif t in aerial:     cls = 'Winged beast'
        elif t in primal:     cls = 'Beast (land)'
        elif t in cav:        cls = 'Cavalry'
        elif t in missile:    cls = 'Missile'
        elif t in inf:        cls = 'Infantry'
        elif t in ELEMENTALS: cls = 'Summoned'
        else:                 cls = 'Infantry'

        # ability text
        if t in ABIL:
            aname, adesc, ahook = ABIL[t]
        elif t in primal:
            d = primal[t]
            sig, cond = PRIMAL_EXTRA.get(t, ('—', '—'))
            sp = (d.get('spells') or [{}])[0]
            lo, hi = (sp.get('dmg') or [0, 0])[:2]
            tgt = 'ONE enemy unit' if sp.get('target') == 'single' else 'EVERY enemy unit'
            mult = None
            # the always-on multiplier, read out of the engine's own desc text
            import re as _re
            mm = _re.search(r'\+(\d+)% strength', d.get('desc', ''))
            if mm: mult = '+' + mm.group(1) + '%'
            aname = '%s + %s (pre-battle strike)' % (sig, sp.get('name', '—'))
            adesc = (
              '%s: %s strength when %s. It is a per-unit multiplier, applied in armStr alongside terrain, '
              'quality and the commander\'s Tactical bonus.%s\n\n'
              '%s (pre-battle strike): in the pre-battle MAGIC phase this unit deals %s–%s strength damage to %s. '
              'Each unit rolls separately to fire, at min(90%%, 30%% + 10%% per Druid level of the commanding character '
              '+ any primalStrike item bonus) — so 30%% with no Druid skill at all, and the 90%% ceiling from Druid 6 up. '
              'A commander whose Druid has been negated (Spiritwrack) contributes nothing.\n\n'
              'Leading it: Druid %s or better, or every missing level costs the unit −20%% strength. A Druidic realm or '
              'a Druidic commander waives that requirement for CREATURE units; Monster units never have it waived that way.'
            ) % (sig, mult or '—', cond,
                 ' Cave Trolls are the one unit with two stacked conditions: ×1.20 always, and a further ×1.30 in mountain '
                 'or forest for about ×1.56 there.' if t == 'Cave Trolls' else '',
                 sp.get('name', '—'), lo, hi, tgt, druid.get(t, 0))
            ahook = 'unitAbilityMult / PRIMAL_ARMY_DEFS_P.spells / primalStrikeChance'
        elif t in ELEMENTALS:
            aname = '— (summoned, temporary)'
            adesc = (
              'Not recruited and not bought: summoned by the Elemental Magic spell Summon Elemental, once per caster per turn. '
              'The engine picks one of the five elementals at random (you cannot choose which), always at VETERAN quality, named '
              '"%s". It joins the caster\'s column immediately and fights as an ordinary army unit with no special ability of its '
              'own — its value is raw strength. It is flagged temporary and DISSOLVES at the end of the turn, leaves no grave, '
              'and can never be raised as a skeleton. It has no cost, no wage, no speed entry and no terrain entry, so terrain '
              'never modifies it (terrainMod falls through to ×1.00).'
            ) % ELEMENTALS[t]
            ahook = 'elemental_summon'
        else:
            aname, adesc, ahook = ('—', 'No special ability recorded in the engine for this unit type.', '—')

        # lift
        if t in CARRIERS:
            lift = '1 hull (2 light foot, or 1 heavy/cavalry/siege/beast). Naval 8 commander: ×2'
        elif t in flying:
            lift = '1 foot unit each'
        elif t in sea:
            lift = 'swims itself — carries nothing'
        elif t in aerial:
            lift = 'flies itself — carries nothing'
        else:
            lift = ''

        # what it costs to ship
        if t in sea or t in aerial or t in flying or t in ELEMENTALS:
            berth = 'n/a'
        elif is_light(t):
            berth = '½ hull'
        else:
            berth = '1 hull (cannot be air-lifted)' if t in cav or t in primal else '1 hull'

        rec = recruit.get(t)
        if t in RAISED_BY_OVERRIDE:
            where = RAISED_BY_OVERRIDE[t][1]
        elif t in ELEMENTALS:
            where = 'Summoned (Summon Elemental) — never recruited'
        elif t in primal:
            where = 'A wild Primal province that offers this beast (any realm, with the Druid level)'
        elif rec:
            where = ', '.join(sorted(rec))
        else:
            where = 'Not in the recruitment table — see notes'

        rows.append({
            'type': t, 'race': race, 'kin': kin, 'cls': cls, 'isprimal': t in primal,
            'druid': druid.get(t, 0),
            'cost': cost.get(t), 'speed': speed.get(t),
            'bstr': bstr.get(t),
            'missile': 'Yes' if t in missile else '',
            'fly': 'Yes' if (t in flying or t in aerial) else '',
            'seaunit': 'Yes' if t in sea else '',
            'garrison': 'Yes' if t in garr else '',
            'tamed': 'Yes' if t in tamed else '',
            'lift': lift, 'berth': berth, 'where': where,
            'aname': aname, 'adesc': adesc, 'ahook': ahook,
            'terr': {k: terr[k].get(t) for k in ['plain','desert','forest','jungle','mountain','sea']},
            'pdesc': (primal[t]['desc'] if t in primal else ''),
        })

    rows.sort(key=lambda r: (RACE_ORDER.get(r['race'], 9), -(r['bstr'] or 0), r['type']))
    return rows


def tmod(v, t, terrain, row):
    """Display value for a terrain cell."""
    if v is not None:
        return v
    if row['fly'] == 'Yes' or row['seaunit'] == 'Yes':
        return 1.0          # terrainMod short-circuits flyers and sea units to 1.0
    if row['kin'] == 'Elemental':
        return 1.0          # no table entry -> ??1.0 default
    return 1.0


def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL; cell.font = HDR_FONT
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[row].height = 34
    ws.freeze_panes = ws.cell(row=row + 1, column=3)


def main():
    rows = build_rows()
    wb = Workbook()

    # ══════════════════════════════════════════════════════════ 1. ARMY UNITS
    ws = wb.active; ws.title = 'Army Units'
    HEADERS = [
        ('Order', 7), ('Unit', 21), ('Raised by (realm)', 12), ('Kinship race', 11), ('Class', 13), ('Druid lvl to lead', 9),
        ('Cost to raise (g)', 9), ('Wage per turn (g)', 9), ('Base strength', 9), ('Speed', 7),
        ('Missile?', 8), ('Flying?', 8), ('Sea unit?', 9), ('Garrison?', 9), ('Dragon-tamed?', 10),
        ('T: plain', 8), ('T: desert', 8), ('T: forest', 8), ('T: jungle', 8), ('T: mountain', 9), ('T: sea', 8),
        ('Where it can be raised', 22), ('Lift it provides', 26), ('Berths it needs at sea', 20),
        ('Special ability', 26), ('EXACTLY how the ability works', 80), ('Engine hook', 24),
        ('CHANGES I WANT', 34),
    ]
    for i, (h, w) in enumerate(HEADERS, start=1):
        ws.cell(row=1, column=i, value=h)
        ws.column_dimensions[get_column_letter(i)].width = w

    r = 2
    for i, row in enumerate(rows, start=1):
        vals = [
            i, row['type'], row['race'], row['kin'], row['cls'], row['druid'],
            row['cost'], row['cost'], row['bstr'], row['speed'],
            row['missile'], row['fly'], row['seaunit'], row['garrison'], row['tamed'],
        ] + [tmod(row['terr'][k], row['type'], k, row) for k in ['plain','desert','forest','jungle','mountain','sea']] + [
            row['where'], row['lift'], row['berth'],
            row['aname'], row['adesc'], row['ahook'], '',
        ]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = NAME_FONT if c == 2 else CELL_FONT
            cell.border = BOX
            if c == len(HEADERS):
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                cell.fill = EDIT_FILL
            elif c >= 22:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            else:
                cell.alignment = Alignment(vertical='top', horizontal='center' if 6 <= c <= 21 else 'left')
            if 16 <= c <= 21 and isinstance(v, (int, float)):
                cell.number_format = '0.00'
        ws.row_dimensions[r].height = 150
        r += 1

    style_header(ws, len(HEADERS))
    last = get_column_letter(len(HEADERS))
    tab = Table(displayName='ArmyUnits', ref='A1:%s%d' % (last, r - 1))
    tab.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
    ws.add_table(tab)
    ws.sheet_view.zoomScale = 85

    # ══════════════════════════════════════════════════════════ 2. ABILITY REFERENCE
    ws2 = wb.create_sheet('Abilities')
    H2 = [('Order', 7), ('Ability', 30), ('Unit(s)', 22), ('Raised by', 11), ('Trigger / condition', 34),
          ('Exact magnitude', 34), ('Full description', 95), ('Engine hook', 26), ('CHANGES I WANT', 34)]
    for i, (h, w) in enumerate(H2, start=1):
        ws2.cell(row=1, column=i, value=h)
        ws2.column_dimensions[get_column_letter(i)].width = w

    TRIG = {
     'Shieldwall':'2+ Shieldwall Levies in the force','Ride Them Down':'You WIN, in plain or desert only',
     'Massed Volley':'Other missile units in the force','Shock of the Charge':'Your side is attacking',
     'Broadside':'Sea battle','Cut the Line':'≥half your winning force are Corsairs',
     'Woodland Volley':'Forest or jungle','Ambush':'Facing Wardens in forest or jungle (either side, any battle)',
     'Loosed From Cover':'Any other unit is still in the force','Skyward Retreat':'Beaten with no retreat',
     'Sea Archers':'A galley anywhere in the force, at sea','Breach and Swamp':'Enemy has sea units',
     'Shield Fort':'≥half the beaten force are Dwarven Guard','Punch Through':'Ironwood Crossbows firing',
     'Bombardment':'Always (order) / vs Garrison units (damage)','Flak Screen':'Enemy fields winged riders',
     'Armoured Hulls':'Always','Run Them Down':'≥half your winning force are Wargs',
     'Outnumbered and Uncaring':'Enemy outnumbers you in UNITS','Sheer Numbers':'Other Goblin Hordes in the force',
     'Terror From Above':'Wyverns present on your side','Reaver Crews':'The phase the force lands from the ships',
     'Screech Dive':'Pre-battle magic phase',
    }
    MAG = {
     'Shieldwall':'×1.25, flat (no growth past 2)','Ride Them Down':'+6% enemy loss chance per unit, cap +18%',
     'Massed Volley':'+10% missile damage per other missile unit, cap +30%','Shock of the Charge':'×1.40 attacking, ×1.00 defending',
     'Broadside':'One extra shot per galleon in the sea archery phase','Cut the Line':'Removes every sea province from the enemy retreat list; vs neutrals, suppresses your own rout losses',
     'Woodland Volley':'×2 missile damage','Ambush':'Enemy archery phase cancelled entirely',
     'Loosed From Cover':'Untargetable until they are the only units left','Skyward Retreat':'Skips the 15% cut-down roll',
     'Sea Archers':'Missile units ×0.75 → ×1.00 at sea','Breach and Swamp':'25% per cohort to sink one enemy sea unit outright',
     'Shield Fort':'No flight casualties; no Ride Them Down; no demotion in retreat — in every battle','Punch Through':'Arrow cap 40%→55% ground, 80%→90% winged (pool shared per target)',
     'Bombardment':'Fires first; ×1.5 damage vs Garrisons','Flak Screen':'−30% base strength on one distinct flyer per gyrocopter',
     'Armoured Hulls':'Immune to all pre-battle missile damage','Run Them Down':'Auto-pursuit (if the loser lives); or ×0.5 own casualties vs neutrals',
     'Outnumbered and Uncaring':'+10% per full 25% excess, cap +30%','Sheer Numbers':'+10% per other horde, cap +30%',
     'Terror From Above':'−15pp enemy fanaticism; ×0.90 neutral defence bonus','Reaver Crews':'×1.25 to the WHOLE force (attacking side only)',
     'Screech Dive':'2–4 damage to one unit, min(90%, 30%+10%/Druid)',
    }

    r2 = 2
    seen = {}
    for row in rows:
        base = row['aname'].split(' + ')[0].split(' (')[0]
        key = row['aname']
        if key in ('—', '— (no combat ability)', '— (summoned, temporary)'):
            continue
        seen.setdefault(key, []).append(row)

    order_n = 1
    for key, group in seen.items():
        row = group[0]
        base = key.split(' + ')[0].split(' (')[0]
        pd = row['pdesc']
        trig = TRIG.get(base)
        mag = MAG.get(base)
        if trig is None and row['isprimal']:
            cond = PRIMAL_EXTRA.get(row['type'], ('', ''))[1]
            trig = 'Always-on: %s; strike: pre-battle magic phase' % cond
        if mag is None and row['isprimal']:
            import re as _re
            mm = _re.search(r'\+(\d+)% strength', pd or '')
            d = RAW['primalDefs'][row['type']]
            sp = (d.get('spells') or [{}])[0]
            lo, hi = (sp.get('dmg') or [0, 0])[:2]
            tgt = 'one unit' if sp.get('target') == 'single' else 'EVERY enemy unit'
            always = ('×1.%s strength' % mm.group(1)) if mm else '—'
            if row['type'] == 'Cave Trolls':
                always = '×1.20 always, ×1.56 in mountain or forest'
            mag = '%s; strike %s–%s to %s at min(90%%, 30%%+10%%/Druid)' % (always, lo, hi, tgt)
        vals = [order_n, key, ', '.join(g['type'] for g in group), row['race'],
                trig or '—', mag or '—', row['adesc'], row['ahook'], '']
        for c, v in enumerate(vals, start=1):
            cell = ws2.cell(row=r2, column=c, value=v)
            cell.font = NAME_FONT if c == 2 else CELL_FONT
            cell.border = BOX
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if c == len(H2): cell.fill = EDIT_FILL
        ws2.row_dimensions[r2].height = 150
        r2 += 1; order_n += 1

    style_header(ws2, len(H2))
    tab2 = Table(displayName='Abilities', ref='A1:%s%d' % (get_column_letter(len(H2)), r2 - 1))
    tab2.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
    ws2.add_table(tab2)
    ws2.sheet_view.zoomScale = 85

    # ══════════════════════════════════════════════════════════ 3. TERRAIN GRID
    ws3 = wb.create_sheet('Terrain grid')
    H3 = ['Order','Unit','Raised by','plain','desert','forest','jungle','mountain','sea','Notes','CHANGES I WANT']
    for i, h in enumerate(H3, start=1):
        ws3.cell(row=1, column=i, value=h)
    for i, w in enumerate([7,21,10,9,9,9,9,10,9,42,34], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    r3 = 2
    for i, row in enumerate(rows, start=1):
        note = ''
        if row['fly'] == 'Yes': note = 'Flyer — terrainMod short-circuits to ×1.00 everywhere.'
        elif row['seaunit'] == 'Yes': note = 'Sea unit — ×1.00 on land and sea alike (the old beached-fleet penalty is gone).'
        elif row['kin'] == 'Elemental': note = 'No table entry — terrainMod default ×1.00.'
        vals = [i, row['type'], row['race']] + [tmod(row['terr'][k], row['type'], k, row) for k in
                ['plain','desert','forest','jungle','mountain','sea']] + [note, '']
        for c, v in enumerate(vals, start=1):
            cell = ws3.cell(row=r3, column=c, value=v)
            cell.font = NAME_FONT if c == 2 else CELL_FONT
            cell.border = BOX
            if c == len(H3): cell.fill = EDIT_FILL; cell.alignment = Alignment(wrap_text=True, vertical='top')
            elif c == 10: cell.alignment = Alignment(wrap_text=True, vertical='top')
            else: cell.alignment = Alignment(vertical='top', horizontal='center' if c >= 4 else 'left')
            if 4 <= c <= 9 and isinstance(v, (int, float)): cell.number_format = '0.00'
        r3 += 1
    style_header(ws3, len(H3))
    tab3 = Table(displayName='TerrainGrid', ref='A1:%s%d' % (get_column_letter(len(H3)), r3 - 1))
    tab3.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
    ws3.add_table(tab3)

    # ══════════════════════════════════════════════════════════ 4. READ ME
    ws4 = wb.create_sheet('Read me', 0)
    ws4.column_dimensions['A'].width = 4
    ws4.column_dimensions['B'].width = 30
    ws4.column_dimensions['C'].width = 108
    lines = [
        ('H', 'Kingdoms Call — Army Units', ''),
        ('', 'Taken from build', 'GM Portal v%s  /  data file v%s' % (STAMP, RAW['dataStamp'])),
        ('', 'Units listed', '%d — 27 humanoid (Human 8, Elven 7, Dwarven 6, Orcish 6), 12 primal, 5 summoned elementals' % len(rows)),
        ('', 'Where the numbers come from',
         'Every figure was read out of the running engine (ARMY_COST_MASTER, ARMY_SPEED_MASTER, UNIT_BASE_STR, '
         'TERRAIN_MODS, DRUID_ARMY_LEVELS_P, PROVINCE_ARMIES_P, PRIMAL_ARMY_DEFS_P), not typed by hand. '
         'The ability prose was written against the engine source; the "Engine hook" column names the function '
         'to look in.'),
        ('', '', ''),
        ('H', 'How to use it', ''),
        ('', 'The yellow column', 'Every sheet ends with a yellow CHANGES I WANT column. Write whatever you want changed there — '
         'plain words are fine ("cost 3 not 4", "drop the mountain bonus", "make this one missile"). Send the file back '
         'and I will apply them to the engine.'),
        ('', 'Sorting', 'Each sheet is a real Excel table, so every header has a sort/filter arrow — click it to sort by cost, '
         'strength, race, anything. The header row and the first two columns stay put when you scroll.'),
        ('', 'Reordering', 'Column A is Order. Retype the numbers however you like (1, 2, 3… or 10, 20, 30 to leave gaps) and '
         'then sort by Order. There are no formulas anywhere in this workbook, so sorting and reordering can never '
         'break a calculation.'),
        ('', 'Cost and wage', 'These are ONE number in the engine: a unit costs its price again every turn it exists. Both columns '
         'are shown so the wage bill is visible, but changing one means changing the other — say which you mean.'),
        ('', '', ''),
        ('H', 'Things worth knowing before you change numbers', ''),
        ('', 'Strength is not the whole story', 'Printed base strength is multiplied by quality (Green ×1, Average ×2, Veteran ×3, '
         'Crack ×4, Elite ×5) before anything else, so a Green Elder Dragon (12) is weaker than an Elite Goblin Horde (15).'),
        ('', 'Terrain × ability × racial all stack', 'They multiply. An Elven Bowmen unit of an Elven realm in forest is ×1.15 terrain '
         '× ×1.5 Elven Woodland Mastery, and its arrows do ×2 damage on top.'),
        ('', 'Racial rules sit outside this sheet', 'Elven Woodland Mastery (×1.5 forest/jungle), Dwarven Bulwark (×1.5 defending), '
         'Human Drillmasters (×2 promotion chance), Orcish war-spoils (wages halved after an attacking win) and the '
         'Elven Homeward Song (a slain Elven unit’s 30/45/60% chance to reach home instead of dying) apply to every '
         'unit of that race, not to one unit type. Say so if you want those changed too.'),
        ('', 'Front-line width caps how many fight', 'plain/desert/sea 6 · forest/jungle 4 · mountain 3. Units past the front line fight '
         'at ×0.5, ×0.25, ×0.125 by line. Piling more units into mountains does not help much.'),
        ('', 'Land units at sea', '×0.75, unless an Elven Galleys is present (missile units only) or the commander has Naval 8.'),
        ('', 'Arrow softening has a ceiling', 'Archery can only reduce a unit to 40% of base strength (80% for a winged rider); '
         'Ironwood Crossbows raise that to 55% / 90%. Dwarf Ironclads are immune outright.'),
        ('', '', ''),
        ('H', 'Found and fixed while building this', ''),
        ('', 'Ironwood Crossbows blurb', 'The player portal’s pop-up said "45% instead of the usual 30%". The engine has said 55% / 40% '
         'since 2026-09-24. Fixed in this build.'),
        ('', 'Primal strike chance blurb', 'The player portal comment said "20% base, +10% per level". The engine uses 30% base. Fixed in this build.'),
        ('', 'Flak Screen never forgot', 'The "already engaged" mark a Gyrocopter puts on an enemy flyer was set and never cleared — not by any of '
         'the three battle-end sweeps, not by the snapshot copier — so it was saved with the game and a winged rider harried '
         'once was immune to Flak Screen for the rest of the GAME. Fixed in this build: the mark now dies with the battle.'),
        ('', '', ''),
        ('', 'Ambush was one-way vs neutrals', 'Only the attacker’s volley was gated in battles against neutral defenders, so Forest Wardens '
         'marching on a neutral garrison did not silence its archers, while the same Wardens attacking a PLAYER did. '
         'Fixed on your ruling, 2026-09-27 — both paths now agree, on attack and on defence.'),
        ('', 'Shield Fort was half-wired vs neutrals', 'The good-order half (immune to Ride Them Down, and no demotion in retreat) was wired in '
         'player-vs-player battles only. Fixed on your ruling, 2026-09-27 — it now works in every battle, and a Guard-heavy '
         'NEUTRAL garrison gets the same protection when it loses.'),
    ]
    rr = 2
    for kind, label, text in lines:
        if kind == 'H':
            c = ws4.cell(row=rr, column=2, value=label)
            c.font = Font(name=FONT, bold=True, size=13, color='2F3E2F')
            ws4.cell(row=rr, column=2).fill = GRP_FILL
            ws4.cell(row=rr, column=3).fill = GRP_FILL
            ws4.row_dimensions[rr].height = 22
        else:
            b = ws4.cell(row=rr, column=2, value=label); b.font = Font(name=FONT, bold=True, size=10)
            b.alignment = Alignment(vertical='top', wrap_text=True)
            t = ws4.cell(row=rr, column=3, value=text); t.font = CELL_FONT
            t.alignment = Alignment(vertical='top', wrap_text=True)
            if text:
                ws4.row_dimensions[rr].height = max(15, 13 * (len(text) // 105 + 1))
        rr += 1
    ws4.sheet_view.showGridLines = False

    out = 'Kingdoms-Call-Army-Units-v%s.xlsx' % STAMP
    wb.save(out)
    print('wrote %s' % out)
    print('  Army Units rows: %d' % (r - 2))
    print('  Abilities rows:  %d' % (r2 - 2))
    print('  Terrain rows:    %d' % (r3 - 2))
    byrace = {}
    for row in rows: byrace[row['race']] = byrace.get(row['race'], 0) + 1
    for k in sorted(byrace, key=lambda x: RACE_ORDER.get(x, 9)):
        print('    %-9s %d' % (k, byrace[k]))


if __name__ == '__main__':
    main()
